from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence


DEFAULT_ELEXIS_SHEETS = ("SrWSD-v2",)
REQUIRED_COLUMNS = ("lemma", "upos", "senseid", "definition")


@dataclass(frozen=True)
class ElexisSense:
    sense_id: str
    lemma: str
    upos: str
    definition: str


class ElexisSenseIndex:
    def __init__(self, senses: Iterable[ElexisSense]):
        self.senses = list(senses)
        self._by_token: dict[tuple[str, str], list[ElexisSense]] = {}
        self._by_lemma: dict[str, list[ElexisSense]] = {}
        seen_by_token: dict[tuple[str, str], set[str]] = {}
        seen_by_lemma: dict[str, set[str]] = {}
        for sense in self.senses:
            token_key = (sense.lemma, sense.upos)
            if sense.sense_id not in seen_by_token.setdefault(token_key, set()):
                self._by_token.setdefault(token_key, []).append(sense)
                seen_by_token[token_key].add(sense.sense_id)
            if sense.sense_id not in seen_by_lemma.setdefault(sense.lemma, set()):
                self._by_lemma.setdefault(sense.lemma, []).append(sense)
                seen_by_lemma[sense.lemma].add(sense.sense_id)

    def candidates_for_token(self, lemma: str, upos: str) -> list[ElexisSense]:
        return list(self._by_token.get((lemma, upos), []))

    def candidates_for_mwe(self, lemma: str) -> list[ElexisSense]:
        return list(self._by_lemma.get(lemma, []))


def load_elexis_senses(
    path: str | Path,
    *,
    sheets: Sequence[str] | None = None,
) -> list[ElexisSense]:
    from openpyxl import load_workbook

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    sheet_names = tuple(sheets or DEFAULT_ELEXIS_SHEETS)
    missing_sheets = [name for name in sheet_names if name not in workbook.sheetnames]
    if missing_sheets:
        raise ValueError(f"Missing Elexis sense sheet(s): {missing_sheets}")

    senses: list[ElexisSense] = []
    seen_rows: set[tuple[str, str, str]] = set()
    for sheet_name in sheet_names:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            continue
        columns = _column_indexes(header)
        if any(column not in columns for column in REQUIRED_COLUMNS):
            continue

        for row in rows:
            sense_id = _cell(row, columns["senseid"])
            lemma = _cell(row, columns["lemma"])
            upos = _cell(row, columns["upos"])
            definition = _cell(row, columns["definition"])
            key = (lemma, upos, sense_id)
            if not sense_id or not lemma or not upos or not definition or key in seen_rows:
                continue
            senses.append(
                ElexisSense(
                    sense_id=sense_id,
                    lemma=lemma,
                    upos=upos,
                    definition=definition,
                )
            )
            seen_rows.add(key)
    return senses


def _column_indexes(header: tuple[object, ...]) -> dict[str, int]:
    return {
        str(value).strip().lower(): index
        for index, value in enumerate(header)
        if value is not None and str(value).strip()
    }


def _cell(row: tuple[object, ...], index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()
