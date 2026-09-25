from __future__ import annotations

import re
import json
import math
import random
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable

from .constants import CONTENT_UPOS, LABELS
from .csvio import read_delimited, write_dicts
from .distilled_wsd import run_distilled_wsd_from_paths
from .evaluation import evaluate_scores, normalize_gold_label
from .scoring import load_polarity_table, load_sentiment_table
from .scoring import label_lemma_score, label_synset_score, score_dataset


ENG30_WORDNET_RE = re.compile(r"^ENG30-\d{8}-[nvars]$")
CUSTOM_PREFIXES = ("CVMWE-", "LLM-", "DRJ0-", "DRS0-", "RSSJ-", "CVDf-", "CV")
REPORT_CONFIGS = [
    ("L0", "lemma", "L0"),
    ("S0", "wsd", "S0-wsd"),
    *[
        item
        for index in range(1, 8)
        for item in (
            (f"S{index}", "wsd", f"A{index}-wsd"),
            (f"S{index}", "avg", f"A{index}-avg2"),
        )
    ],
]
DEFAULT_SIGNIFICANCE_COMPARISONS = [
    ("S7", "avg", "L0", "lemma"),
    ("S7", "avg", "S0", "wsd"),
    ("S7", "wsd", "S0", "wsd"),
    ("S7", "avg", "S7", "wsd"),
    ("S7", "avg", "S6", "avg"),
]


def classify_sense_id(sense_id: str) -> str:
    if sense_id == "NEW_SENSE":
        return "new_sense"
    if ENG30_WORDNET_RE.match(sense_id):
        return "eng30_wordnet_like"
    if sense_id.startswith("ENG30-"):
        return "eng30_nonstandard"
    if sense_id.startswith(CUSTOM_PREFIXES):
        return "elexis_custom_or_mwe"
    return "other"


def build_missing_sentiment_reports(
    wsd_rows: list[dict[str, str]],
    sentiment_ids_by_table: dict[str, set[str]],
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    mentions: dict[str, dict[str, object]] = {}
    order: list[str] = []
    for row in wsd_rows:
        kbid = (row.get("KBid") or "").strip()
        if not kbid or kbid == "NEW_SENSE":
            continue
        if kbid not in mentions:
            mentions[kbid] = {
                "count": 0,
                "sentences": set(),
                "lemmas": set(),
                "target_types": set(),
                "upos": set(),
            }
            order.append(kbid)
        entry = mentions[kbid]
        entry["count"] = int(entry["count"]) + 1
        _add_nonblank(entry["sentences"], row.get("sentence_id"))
        _add_nonblank(entry["lemmas"], row.get("lemma"))
        _add_nonblank(entry["target_types"], row.get("target_type"))
        _add_nonblank(entry["upos"], row.get("upos"))

    by_id: list[dict[str, object]] = []
    class_counts: dict[tuple[str, str], dict[str, object]] = {}
    for table in sorted(sentiment_ids_by_table):
        sentiment_ids = sentiment_ids_by_table[table]
        for kbid in order:
            if kbid in sentiment_ids:
                continue
            entry = mentions[kbid]
            id_class = classify_sense_id(kbid)
            by_id.append(
                {
                    "table": table,
                    "KBid": kbid,
                    "id_class": id_class,
                    "prefix": sense_id_prefix(kbid),
                    "count": entry["count"],
                    "sentences": _join_limited(entry["sentences"]),
                    "lemmas": _join_limited(entry["lemmas"]),
                    "target_types": _join_limited(entry["target_types"]),
                    "upos": _join_limited(entry["upos"]),
                }
            )
            key = (table, id_class)
            bucket = class_counts.setdefault(key, {"ids": 0, "mentions": 0})
            bucket["ids"] = int(bucket["ids"]) + 1
            bucket["mentions"] = int(bucket["mentions"]) + int(entry["count"])

    by_class = [
        {"table": table, "id_class": id_class, **counts}
        for (table, id_class), counts in sorted(class_counts.items())
    ]
    return by_id, by_class


def build_lemma_wsd_overlap_rows(
    token_rows: list[dict[str, str]],
    wsd_rows: list[dict[str, str]],
    polarity: dict[str, int],
    sentiment_ids: set[str],
) -> list[dict[str, object]]:
    wsd_by_token = _wsd_by_single_token(wsd_rows)
    rows: list[dict[str, object]] = []
    for token in token_rows:
        if token.get("upos") not in CONTENT_UPOS:
            continue
        lemma = token.get("lemma", "")
        if not lemma or lemma[0].isdigit():
            continue
        key = (token.get("sample_id", ""), token.get("token_index", ""))
        decision = wsd_by_token.get(key)
        kbid = decision.get("KBid", "") if decision else ""
        lemma_covered = lemma in polarity
        wsd_covered = bool(kbid and kbid in sentiment_ids)
        rows.append(
            {
                "sample_id": token.get("sample_id", ""),
                "token_index": token.get("token_index", ""),
                "lemma": lemma,
                "upos": token.get("upos", ""),
                "KBid": kbid,
                "id_class": classify_sense_id(kbid) if kbid else "",
                "lemma_has_polarity": str(lemma_covered),
                "wsd_has_sentiment": str(wsd_covered),
                "quadrant": _overlap_quadrant(lemma_covered, wsd_covered),
            }
        )
    return rows


def build_wsd_stage_coverage_rows(
    wsd_rows: list[dict[str, str]],
    sentiment_ids: set[str],
) -> list[dict[str, object]]:
    groups: dict[str, list[dict[str, str]]] = {"all": wsd_rows}
    for row in wsd_rows:
        groups.setdefault(f"target_type={row.get('target_type', '')}", []).append(row)
        groups.setdefault(f"upos={row.get('upos', '')}", []).append(row)

    return [
        {"group": name, **_stage_counts(rows, sentiment_ids)}
        for name, rows in groups.items()
    ]


def build_prediction_wide_rows(score_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    by_sample: dict[str, dict[str, str]] = {}
    config_order: list[str] = []
    for row in score_rows:
        config = row.get("config", "")
        if config and config not in config_order:
            config_order.append(config)
        sample_id = row.get("sample_id", "")
        base = by_sample.setdefault(
            sample_id,
            {
                "sample_id": sample_id,
                "sentence_text": row.get("sentence_text", ""),
                "gold_label": row.get("gold_label", ""),
            },
        )
        base[f"{config}_label"] = row.get("label", "")
        base[f"{config}_score"] = row.get("score", "")
        base[f"{config}_coverage"] = row.get("coverage", "")

    rows: list[dict[str, str]] = []
    for sample_id in by_sample:
        row = {
            "sample_id": by_sample[sample_id]["sample_id"],
            "sentence_text": by_sample[sample_id]["sentence_text"],
            "gold_label": by_sample[sample_id]["gold_label"],
        }
        for config in config_order:
            for suffix in ("label", "score", "coverage"):
                key = f"{config}_{suffix}"
                if key in by_sample[sample_id]:
                    row[key] = by_sample[sample_id][key]
        rows.append(row)
    return rows


def sense_id_prefix(sense_id: str) -> str:
    if "-" in sense_id:
        return sense_id.split("-", 1)[0]
    match = re.match(r"^[A-Za-z]+\d*", sense_id)
    return match.group(0) if match else sense_id


def load_sentiment_ids_by_table(resources_dir: str | Path) -> dict[str, set[str]]:
    synset_dir = Path(resources_dir) / "synset_sentiment"
    return {
        path.stem: set(load_sentiment_table(path))
        for path in sorted(synset_dir.glob("*.csv"))
    }


def write_analysis_tables(
    *,
    token_path: str | Path,
    wsd_path: str | Path,
    scores_path: str | Path,
    resources_dir: str | Path,
    out_dir: str | Path,
) -> dict[str, Path]:
    out_path = Path(out_dir)
    out_path.mkdir(parents=True, exist_ok=True)

    token_rows = read_delimited(token_path, delimiter="\t")
    wsd_rows = read_delimited(wsd_path, delimiter="\t")
    score_rows = read_delimited(scores_path, delimiter="\t")
    sentiment_ids_by_table = load_sentiment_ids_by_table(resources_dir)
    sentiment_ids = set().union(*sentiment_ids_by_table.values()) if sentiment_ids_by_table else set()
    polarity = load_polarity_table(Path(resources_dir) / "lemma_polarity.csv")

    missing_by_id, missing_by_class = build_missing_sentiment_reports(wsd_rows, sentiment_ids_by_table)
    overlap_rows = build_lemma_wsd_overlap_rows(token_rows, wsd_rows, polarity, sentiment_ids)
    stage_rows = build_wsd_stage_coverage_rows(wsd_rows, sentiment_ids)
    prediction_rows = build_prediction_wide_rows(score_rows)

    paths = {
        "missing_by_id": out_path / "missing_sentiment_by_id.tsv",
        "missing_by_class": out_path / "missing_sentiment_by_class.tsv",
        "overlap": out_path / "lemma_wsd_overlap.tsv",
        "stage": out_path / "wsd_stage_coverage.tsv",
        "prediction_wide": out_path / "prediction_wide.tsv",
    }
    write_dicts(paths["missing_by_id"], missing_by_id, _fieldnames(missing_by_id), delimiter="\t")
    write_dicts(paths["missing_by_class"], missing_by_class, _fieldnames(missing_by_class), delimiter="\t")
    write_dicts(paths["overlap"], overlap_rows, _fieldnames(overlap_rows), delimiter="\t")
    write_dicts(paths["stage"], stage_rows, _fieldnames(stage_rows), delimiter="\t")
    write_dicts(paths["prediction_wide"], prediction_rows, _fieldnames(prediction_rows), delimiter="\t")
    return paths


def run_full_analysis(
    *,
    gold_workbook: str | Path,
    sample: str | Path,
    tokens: str | Path,
    annotations_jsonl: str | Path,
    sense_repo: str | Path,
    model: str | Path,
    resources: str | Path,
    out_dir: str | Path,
    theta: float = 0.33,
    theta_sweep: list[float] | None = None,
    reuse_wsd: str | Path | None = None,
    sense_sheets: list[str] | None = None,
    text_prefix: str | None = None,
) -> dict[str, object]:
    out_path = Path(out_dir)
    out_path.mkdir(parents=True, exist_ok=True)

    gold_rows = load_gold_rows(gold_workbook)
    gold_sample_path = out_path / "gold_sample.tsv"
    gold_tokens_path = out_path / "gold_tokens.tsv"
    gold_annotations_path = out_path / "gold_annotations.jsonl"
    write_gold_subset_files(
        sample_path=sample,
        token_path=tokens,
        annotations_jsonl=annotations_jsonl,
        gold_rows=gold_rows,
        sample_out=gold_sample_path,
        token_out=gold_tokens_path,
        annotations_out=gold_annotations_path,
    )

    wsd_path = Path(reuse_wsd) if reuse_wsd else out_path / "wsd_distilled_full.tsv"
    if reuse_wsd is None:
        run_distilled_wsd_from_paths(
            annotations_jsonl=annotations_jsonl,
            sense_repo=sense_repo,
            model=model,
            out_path=wsd_path,
            sense_sheets=sense_sheets,
            text_prefix=text_prefix,
        )

    full_scores_path = out_path / "scores_full.tsv"
    score_dataset(
        sample_path=sample,
        token_path=tokens,
        wsd_path=wsd_path,
        resources_dir=resources,
        out_path=full_scores_path,
        theta=theta,
    )
    gold_scores_path = out_path / "scores_gold.tsv"
    score_dataset(
        sample_path=gold_sample_path,
        token_path=gold_tokens_path,
        wsd_path=wsd_path,
        resources_dir=resources,
        out_path=gold_scores_path,
        theta=theta,
    )
    evaluation_dir = out_path / "evaluation"
    evaluation_report = evaluate_scores(scores_path=gold_scores_path, out_dir=evaluation_dir)

    metrics_path = out_path / "metrics_by_config.tsv"
    metrics_rows = read_delimited(evaluation_dir / "metrics.csv", delimiter=",")
    write_dicts(metrics_path, metrics_rows, ["config", "n", "accuracy", "macro_f1", "coverage_mean"], delimiter="\t")

    analysis_paths = write_analysis_tables(
        token_path=tokens,
        wsd_path=wsd_path,
        scores_path=gold_scores_path,
        resources_dir=resources,
        out_dir=out_path,
    )
    theta_sweep_path = write_theta_sweep(
        gold_sample_path=gold_sample_path,
        gold_tokens_path=gold_tokens_path,
        wsd_path=wsd_path,
        resources_dir=resources,
        out_dir=out_path,
        theta_values=theta_sweep or [0.0, 0.1, 0.2, 0.33, 0.5],
    )
    report_paths = write_report_outputs(
        scores_path=gold_scores_path,
        theta_sweep_path=theta_sweep_path,
        out_dir=out_path,
    )
    summary_path = write_analysis_summary(
        out_path / "analysis_summary.md",
        gold_rows=gold_rows,
        evaluation_report=evaluation_report,
        metrics_rows=metrics_rows,
        missing_by_class=read_delimited(analysis_paths["missing_by_class"], delimiter="\t"),
        stage_rows=read_delimited(analysis_paths["stage"], delimiter="\t"),
        report_metrics=read_delimited(report_paths["report_metrics"], delimiter="\t"),
        significance_rows=read_delimited(report_paths["significance_tests"], delimiter="\t"),
    )

    result = {
        "out_dir": str(out_path),
        "gold_rows": len(gold_rows),
        "wsd_path": str(wsd_path),
        "full_scores": str(full_scores_path),
        "gold_scores": str(gold_scores_path),
        "metrics_by_config": str(metrics_path),
        "theta_sweep": str(theta_sweep_path),
        "analysis_summary": str(summary_path),
        **{name: str(path) for name, path in analysis_paths.items()},
        **{name: str(path) for name, path in report_paths.items()},
    }
    (out_path / "run_manifest.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return result


def load_gold_rows(path: str | Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Combined"] if "Combined" in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []
    headers = [str(value or "") for value in header]
    records: list[dict[str, str]] = []
    seen: set[str] = set()
    for values in rows:
        row = {
            header: "" if value is None else str(value)
            for header, value in zip(headers, values)
        }
        sample_id = row.get("sample_id", "").strip()
        normalized = normalize_gold_label(row.get("sentiment_label"))
        if not sample_id or not normalized or sample_id in seen:
            continue
        row["sentiment_label"] = row.get("sentiment_label", "").strip()
        row["normalized_gold"] = normalized
        records.append(row)
        seen.add(sample_id)
    return records


def write_gold_subset_files(
    *,
    sample_path: str | Path,
    token_path: str | Path,
    annotations_jsonl: str | Path,
    gold_rows: list[dict[str, str]],
    sample_out: str | Path,
    token_out: str | Path,
    annotations_out: str | Path,
) -> None:
    import csv

    gold_by_id = {row["sample_id"]: row for row in gold_rows}
    gold_ids = set(gold_by_id)

    with Path(sample_path).open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        fieldnames = reader.fieldnames or []
        sample_rows = []
        for row in reader:
            sample_id = row.get("sample_id", "")
            if sample_id not in gold_ids:
                continue
            row["sentiment_label"] = gold_by_id[sample_id]["sentiment_label"]
            if "annotator" in row and gold_by_id[sample_id].get("annotator"):
                row["annotator"] = gold_by_id[sample_id]["annotator"]
            sample_rows.append(row)

    ordered_samples = {row["sample_id"]: row for row in sample_rows}
    sample_rows = [ordered_samples[row["sample_id"]] for row in gold_rows if row["sample_id"] in ordered_samples]
    with Path(sample_out).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter="\t")
        writer.writeheader()
        writer.writerows(sample_rows)

    with Path(token_path).open("r", encoding="utf-8-sig", newline="") as src, Path(token_out).open(
        "w", encoding="utf-8-sig", newline=""
    ) as dst:
        reader = csv.DictReader(src, delimiter="\t")
        writer = csv.DictWriter(dst, fieldnames=reader.fieldnames or [], delimiter="\t")
        writer.writeheader()
        for row in reader:
            if row.get("sample_id") in gold_ids:
                writer.writerow(row)

    with Path(annotations_jsonl).open("r", encoding="utf-8") as src, Path(annotations_out).open(
        "w", encoding="utf-8", newline="\n"
    ) as dst:
        for line in src:
            if not line.strip():
                continue
            payload = json.loads(line)
            if payload.get("id") in gold_ids:
                dst.write(json.dumps(payload, ensure_ascii=False) + "\n")


def write_theta_sweep(
    *,
    gold_sample_path: str | Path,
    gold_tokens_path: str | Path,
    wsd_path: str | Path,
    resources_dir: str | Path,
    out_dir: str | Path,
    theta_values: list[float],
) -> Path:
    rows: list[dict[str, object]] = []
    out_path = Path(out_dir)
    for theta in theta_values:
        theta_label = str(theta).replace(".", "p")
        scores_path = out_path / f"scores_gold_theta_{theta_label}.tsv"
        eval_dir = out_path / f"evaluation_theta_{theta_label}"
        score_dataset(
            sample_path=gold_sample_path,
            token_path=gold_tokens_path,
            wsd_path=wsd_path,
            resources_dir=resources_dir,
            out_path=scores_path,
            theta=theta,
        )
        evaluate_scores(scores_path=scores_path, out_dir=eval_dir)
        for metric in read_delimited(eval_dir / "metrics.csv", delimiter=","):
            rows.append({"theta": theta, **metric})
    path = out_path / "theta_sweep.tsv"
    write_dicts(path, rows, ["theta", "config", "n", "accuracy", "macro_f1", "coverage_mean"], delimiter="\t")
    return path


def write_report_outputs(
    *,
    scores_path: str | Path,
    theta_sweep_path: str | Path,
    out_dir: str | Path,
) -> dict[str, Path]:
    out_path = Path(out_dir)
    theta_rows = read_delimited(theta_sweep_path, delimiter="\t")
    score_rows = read_delimited(scores_path, delimiter="\t")
    report_metrics = build_report_metric_rows(theta_rows)
    comparison_rows = build_wsd_avg_comparison_rows(report_metrics)
    significance_rows = build_significance_rows(score_rows, report_metrics)

    paths = {
        "report_metrics": out_path / "report_metrics.tsv",
        "report_wsd_avg_comparison": out_path / "report_wsd_avg_comparison.tsv",
        "significance_tests": out_path / "significance_tests.tsv",
    }
    write_dicts(
        paths["report_metrics"],
        report_metrics,
        ["resource", "method", "internal_config", "theta", "n", "accuracy", "macro_f1", "coverage_mean"],
        delimiter="\t",
    )
    write_dicts(
        paths["report_wsd_avg_comparison"],
        comparison_rows,
        [
            "resource",
            "wsd_theta",
            "wsd_accuracy",
            "wsd_macro_f1",
            "wsd_coverage",
            "avg_theta",
            "avg_accuracy",
            "avg_macro_f1",
            "avg_coverage",
            "delta_avg_minus_wsd",
        ],
        delimiter="\t",
    )
    write_dicts(
        paths["significance_tests"],
        significance_rows,
        [
            "comparison",
            "system_a",
            "system_b",
            "n",
            "accuracy_a",
            "accuracy_b",
            "delta_accuracy",
            "mcnemar_b",
            "mcnemar_c",
            "mcnemar_p",
            "macro_f1_a",
            "macro_f1_b",
            "delta_macro_f1",
            "randomization_iterations",
            "randomization_p",
            "holm_p",
            "significant_0_05",
        ],
        delimiter="\t",
    )
    return paths


def build_report_metric_rows(theta_rows: list[dict[str, str]]) -> list[dict[str, object]]:
    rows_by_config: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in theta_rows:
        rows_by_config[row.get("config", "")].append(row)

    report_rows: list[dict[str, object]] = []
    for resource, method, internal_config in REPORT_CONFIGS:
        candidates = rows_by_config.get(internal_config, [])
        if not candidates:
            continue
        best = max(candidates, key=lambda row: (float(row.get("macro_f1") or 0), float(row.get("accuracy") or 0)))
        report_rows.append(
            {
                "resource": resource,
                "method": method,
                "internal_config": internal_config,
                "theta": best.get("theta", ""),
                "n": best.get("n", ""),
                "accuracy": best.get("accuracy", ""),
                "macro_f1": best.get("macro_f1", ""),
                "coverage_mean": best.get("coverage_mean", ""),
            }
        )
    return report_rows


def build_wsd_avg_comparison_rows(report_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    by_key = {
        (str(row.get("resource", "")), str(row.get("method", ""))): row
        for row in report_rows
    }
    rows: list[dict[str, object]] = []
    for index in range(1, 8):
        resource = f"S{index}"
        wsd = by_key.get((resource, "wsd"))
        avg = by_key.get((resource, "avg"))
        if not wsd or not avg:
            continue
        rows.append(
            {
                "resource": resource,
                "wsd_theta": wsd.get("theta", ""),
                "wsd_accuracy": wsd.get("accuracy", ""),
                "wsd_macro_f1": wsd.get("macro_f1", ""),
                "wsd_coverage": wsd.get("coverage_mean", ""),
                "avg_theta": avg.get("theta", ""),
                "avg_accuracy": avg.get("accuracy", ""),
                "avg_macro_f1": avg.get("macro_f1", ""),
                "avg_coverage": avg.get("coverage_mean", ""),
                "delta_avg_minus_wsd": round(float(avg.get("macro_f1") or 0) - float(wsd.get("macro_f1") or 0), 12),
            }
        )
    return rows


def build_significance_rows(
    score_rows: list[dict[str, str]],
    report_rows: list[dict[str, object]],
    *,
    comparisons: list[tuple[str, str, str, str]] | None = None,
    randomization_iterations: int = 2000,
    random_seed: int = 42,
) -> list[dict[str, object]]:
    specs = {
        (str(row.get("resource", "")), str(row.get("method", ""))): row
        for row in report_rows
    }
    predictions = {
        key: _predictions_for_report_config(score_rows, row)
        for key, row in specs.items()
    }
    rows: list[dict[str, object]] = []
    p_values: list[float] = []
    active_comparisons = comparisons or DEFAULT_SIGNIFICANCE_COMPARISONS
    for comparison in active_comparisons:
        left = (comparison[0], comparison[1])
        right = (comparison[2], comparison[3])
        left_predictions = predictions.get(left, {})
        right_predictions = predictions.get(right, {})
        sample_ids = sorted(set(left_predictions) & set(right_predictions))
        if not sample_ids:
            continue
        gold = [left_predictions[sample_id][0] for sample_id in sample_ids]
        pred_a = [left_predictions[sample_id][1] for sample_id in sample_ids]
        pred_b = [right_predictions[sample_id][1] for sample_id in sample_ids]
        accuracy_a = _accuracy(gold, pred_a)
        accuracy_b = _accuracy(gold, pred_b)
        macro_a = _macro_f1(gold, pred_a)
        macro_b = _macro_f1(gold, pred_b)
        mcnemar_b, mcnemar_c = _mcnemar_counts(gold, pred_a, pred_b)
        randomization_p = approximate_randomization_p(
            gold,
            pred_a,
            pred_b,
            iterations=randomization_iterations,
            seed=random_seed,
        )
        p_values.append(randomization_p)
        rows.append(
            {
                "comparison": f"{_system_name(left)} vs {_system_name(right)}",
                "system_a": _system_name(left),
                "system_b": _system_name(right),
                "n": len(sample_ids),
                "accuracy_a": round(accuracy_a, 12),
                "accuracy_b": round(accuracy_b, 12),
                "delta_accuracy": round(accuracy_a - accuracy_b, 12),
                "mcnemar_b": mcnemar_b,
                "mcnemar_c": mcnemar_c,
                "mcnemar_p": exact_mcnemar_p(mcnemar_b, mcnemar_c),
                "macro_f1_a": round(macro_a, 12),
                "macro_f1_b": round(macro_b, 12),
                "delta_macro_f1": round(macro_a - macro_b, 12),
                "randomization_iterations": randomization_iterations,
                "randomization_p": randomization_p,
            }
        )

    adjusted = holm_bonferroni(p_values)
    for row, holm_p in zip(rows, adjusted):
        row["holm_p"] = holm_p
        row["significant_0_05"] = str(holm_p < 0.05)
    return rows


def exact_mcnemar_p(b: int, c: int) -> float:
    n = b + c
    if n == 0:
        return 1.0
    lower_tail = sum(math.comb(n, index) for index in range(0, min(b, c) + 1)) / (2 ** n)
    return round(min(1.0, 2 * lower_tail), 12)


def approximate_randomization_p(
    gold: list[str],
    pred_a: list[str],
    pred_b: list[str],
    *,
    iterations: int,
    seed: int,
) -> float:
    observed = _macro_f1(gold, pred_a) - _macro_f1(gold, pred_b)
    rng = random.Random(seed)
    more_extreme = 0
    for _ in range(iterations):
        shuffled_a: list[str] = []
        shuffled_b: list[str] = []
        for left, right in zip(pred_a, pred_b):
            if rng.random() < 0.5:
                shuffled_a.append(right)
                shuffled_b.append(left)
            else:
                shuffled_a.append(left)
                shuffled_b.append(right)
        delta = _macro_f1(gold, shuffled_a) - _macro_f1(gold, shuffled_b)
        if abs(delta) >= abs(observed) - 1e-12:
            more_extreme += 1
    return round((more_extreme + 1) / (iterations + 1), 12)


def holm_bonferroni(p_values: list[float]) -> list[float]:
    ordered = sorted(enumerate(p_values), key=lambda item: item[1])
    adjusted = [0.0 for _ in p_values]
    previous = 0.0
    total = len(p_values)
    for rank, (original_index, p_value) in enumerate(ordered):
        corrected = min(1.0, (total - rank) * p_value)
        corrected = max(previous, corrected)
        adjusted[original_index] = round(corrected, 12)
        previous = corrected
    return adjusted


def write_analysis_summary(
    path: str | Path,
    *,
    gold_rows: list[dict[str, str]],
    evaluation_report: dict[str, object],
    metrics_rows: list[dict[str, str]],
    missing_by_class: list[dict[str, str]],
    stage_rows: list[dict[str, str]],
    report_metrics: list[dict[str, str]] | None = None,
    significance_rows: list[dict[str, str]] | None = None,
) -> Path:
    top_metrics = sorted(
        metrics_rows,
        key=lambda row: (float(row.get("accuracy") or 0), float(row.get("macro_f1") or 0)),
        reverse=True,
    )[:5]
    lines = [
        "# Full Distilled WSD Analysis",
        "",
        f"Gold sentences: {len(gold_rows)}",
        f"Evaluated score rows: {evaluation_report.get('evaluated_rows', 0)}",
        "",
        "## Top Configs",
        "",
        "| Config | Accuracy | Macro-F1 | Coverage |",
        "| --- | ---: | ---: | ---: |",
    ]
    for row in top_metrics:
        lines.append(f"| {row['config']} | {row['accuracy']} | {row['macro_f1']} | {row['coverage_mean']} |")
    lines.extend(["", "## Missing Sentiment By Class", "", "| Table | ID class | IDs | Mentions |", "| --- | --- | ---: | ---: |"])
    for row in missing_by_class[:20]:
        lines.append(f"| {row['table']} | {row['id_class']} | {row['ids']} | {row['mentions']} |")
    all_stage = next((row for row in stage_rows if row.get("group") == "all"), None)
    if all_stage:
        lines.extend([
            "",
            "## WSD Stage Coverage",
            "",
            f"Eligible targets: {all_stage['eligible_targets']}",
            f"Candidate found: {all_stage['candidates_found']}",
            f"Selected ID has sentiment: {all_stage['selected_has_sentiment']}",
            f"Any candidate has sentiment: {all_stage['candidate_any_has_sentiment']}",
        ])
    if report_metrics:
        lines.extend([
            "",
            "## Report Metrics",
            "",
            "| Resource | Method | Theta | Accuracy | Macro-F1 | Coverage |",
            "| --- | --- | ---: | ---: | ---: | ---: |",
        ])
        for row in sorted(report_metrics, key=lambda item: float(item.get("macro_f1") or 0), reverse=True)[:10]:
            lines.append(
                f"| {row['resource']} | {row['method']} | {row['theta']} | "
                f"{row['accuracy']} | {row['macro_f1']} | {row['coverage_mean']} |"
            )
    if significance_rows:
        lines.extend([
            "",
            "## Main Significance Tests",
            "",
            "| Comparison | Delta Macro-F1 | Randomization p | Holm p | Significant |",
            "| --- | ---: | ---: | ---: | --- |",
        ])
        for row in significance_rows:
            lines.append(
                f"| {row['comparison']} | {row['delta_macro_f1']} | {row['randomization_p']} | "
                f"{row['holm_p']} | {row['significant_0_05']} |"
            )
    out_path = Path(path)
    out_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out_path


def _add_nonblank(values: object, value: object) -> None:
    if value:
        values.add(str(value))


def _join_limited(values: object, limit: int = 25) -> str:
    return ";".join(sorted(str(value) for value in values)[:limit])


def _wsd_by_single_token(wsd_rows: list[dict[str, str]]) -> dict[tuple[str, str], dict[str, str]]:
    result: dict[tuple[str, str], dict[str, str]] = {}
    for row in wsd_rows:
        token_indices = row.get("token_indices", "")
        if "," in token_indices:
            continue
        result[(row.get("sentence_id", ""), token_indices)] = row
    return result


def _overlap_quadrant(lemma_covered: bool, wsd_covered: bool) -> str:
    if lemma_covered and wsd_covered:
        return "both"
    if lemma_covered:
        return "lemma_only"
    if wsd_covered:
        return "wsd_only"
    return "neither"


def _stage_counts(rows: list[dict[str, str]], sentiment_ids: set[str]) -> dict[str, int]:
    return {
        "eligible_targets": len(rows),
        "candidates_found": sum(1 for row in rows if int(row.get("NumberOfSenses") or 0) > 0),
        "distilled": sum(1 for row in rows if row.get("Origine") == "DISTILLED"),
        "first": sum(1 for row in rows if row.get("Origine") == "FIRST"),
        "new_sense": sum(1 for row in rows if row.get("KBid") == "NEW_SENSE" or row.get("Origine") == "None"),
        "selected_has_sentiment": sum(1 for row in rows if row.get("KBid", "") in sentiment_ids),
        "candidate_any_has_sentiment": sum(1 for row in rows if _candidate_any_has_sentiment(row, sentiment_ids)),
    }


def _candidate_any_has_sentiment(row: dict[str, str], sentiment_ids: set[str]) -> bool:
    candidates = [
        candidate.split("[", 1)[0].strip()
        for candidate in (row.get("Possible") or "").split(";")
        if candidate.strip()
    ]
    return any(candidate in sentiment_ids for candidate in candidates)


def _predictions_for_report_config(
    score_rows: list[dict[str, str]],
    report_row: dict[str, object],
) -> dict[str, tuple[str, str]]:
    internal_config = str(report_row.get("internal_config", ""))
    theta = float(report_row.get("theta") or 0)
    predictions: dict[str, tuple[str, str]] = {}
    for row in score_rows:
        if row.get("config") != internal_config:
            continue
        sample_id = row.get("sample_id", "")
        gold = normalize_gold_label(row.get("gold_label"))
        if not sample_id or not gold:
            continue
        score = float(row.get("score") or 0)
        predicted = label_lemma_score(score) if internal_config == "L0" else label_synset_score(score, theta=theta)
        predictions[sample_id] = (gold, predicted)
    return predictions


def _system_name(key: tuple[str, str]) -> str:
    return f"{key[0]}-{key[1]}"


def _accuracy(gold: list[str], predicted: list[str]) -> float:
    if not gold:
        return 0.0
    return sum(1 for expected, actual in zip(gold, predicted) if expected == actual) / len(gold)


def _macro_f1(gold: list[str], predicted: list[str]) -> float:
    if not gold:
        return 0.0
    f1_values = []
    for label in LABELS:
        tp = sum(1 for expected, actual in zip(gold, predicted) if expected == label and actual == label)
        fp = sum(1 for expected, actual in zip(gold, predicted) if expected != label and actual == label)
        fn = sum(1 for expected, actual in zip(gold, predicted) if expected == label and actual != label)
        precision = tp / (tp + fp) if tp + fp else 0.0
        recall = tp / (tp + fn) if tp + fn else 0.0
        f1_values.append(2 * precision * recall / (precision + recall) if precision + recall else 0.0)
    return sum(f1_values) / len(f1_values)


def _mcnemar_counts(gold: list[str], pred_a: list[str], pred_b: list[str]) -> tuple[int, int]:
    a_only = 0
    b_only = 0
    for expected, left, right in zip(gold, pred_a, pred_b):
        left_correct = left == expected
        right_correct = right == expected
        if left_correct and not right_correct:
            a_only += 1
        elif right_correct and not left_correct:
            b_only += 1
    return a_only, b_only


def _fieldnames(rows: list[dict[str, object]]) -> list[str]:
    if rows:
        return list(rows[0])
    return ["empty"]
