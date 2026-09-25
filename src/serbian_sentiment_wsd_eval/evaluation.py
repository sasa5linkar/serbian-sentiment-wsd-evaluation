from __future__ import annotations

import csv
from collections import Counter, defaultdict
from pathlib import Path

from .constants import LABELS
from .csvio import read_delimited, write_dicts


def normalize_gold_label(value: object) -> str | None:
    text = "" if value is None else str(value).strip().lower()
    if not text:
        return None
    mapping = {
        "positive": "positive",
        "pos": "positive",
        "+1": "positive",
        "1": "positive",
        "1.0": "positive",
        "negative": "negative",
        "neg": "negative",
        "-1": "negative",
        "-1.0": "negative",
        "neutral": "neutral",
        "neu": "neutral",
        "0": "neutral",
        "0.0": "neutral",
    }
    return mapping.get(text)


def evaluate_scores(
    *,
    scores_path: str | Path,
    out_dir: str | Path,
    gold_path: str | Path | None = None,
) -> dict[str, object]:
    scores = read_delimited(scores_path, delimiter="\t")
    gold_by_sample = _load_gold(gold_path) if gold_path else {}
    evaluated: list[dict[str, str]] = []

    for row in scores:
        gold = gold_by_sample.get(row.get("sample_id", "")) or normalize_gold_label(row.get("gold_label"))
        if not gold:
            continue
        predicted = normalize_gold_label(row.get("label")) or "neutral"
        evaluated.append({**row, "gold": gold, "predicted": predicted})

    out_path = Path(out_dir)
    out_path.mkdir(parents=True, exist_ok=True)
    metrics = _metrics_by_config(evaluated)
    confusion = _confusion_by_config(evaluated)

    write_dicts(
        out_path / "metrics.csv",
        metrics,
        ["config", "n", "accuracy", "macro_f1", "coverage_mean"],
    )
    write_dicts(
        out_path / "confusion_matrix.csv",
        confusion,
        ["config", "gold", "predicted", "count"],
    )
    _write_summary(out_path / "summary.md", metrics, bool(evaluated), gold_path)
    return {
        "scores_path": str(scores_path),
        "gold_path": str(gold_path) if gold_path else "",
        "evaluated_rows": len(evaluated),
        "metrics_path": str(out_path / "metrics.csv"),
        "confusion_path": str(out_path / "confusion_matrix.csv"),
        "summary_path": str(out_path / "summary.md"),
    }


def _load_gold(path: str | Path | None) -> dict[str, str]:
    if path is None:
        return {}
    gold_path = Path(path)
    if gold_path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _load_gold_xlsx(gold_path)
    delimiter = "\t" if gold_path.suffix.lower() == ".tsv" else None
    rows = read_delimited(gold_path, delimiter=delimiter)
    return _gold_from_rows(rows)


def _load_gold_xlsx(path: Path) -> dict[str, str]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {}
    headers = [str(value or "") for value in rows[0]]
    records: list[dict[str, object]] = []
    for values in rows[1:]:
        records.append({header: value for header, value in zip(headers, values)})
    return _gold_from_rows(records)


def _gold_from_rows(rows: list[dict[str, object]]) -> dict[str, str]:
    gold: dict[str, str] = {}
    for row in rows:
        sample_id = str(row.get("sample_id") or "").strip()
        label = normalize_gold_label(row.get("sentiment_label"))
        if sample_id and label:
            gold[sample_id] = label
    return gold


def _metrics_by_config(rows: list[dict[str, str]]) -> list[dict[str, object]]:
    by_config: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        by_config[row.get("config", "")].append(row)

    metrics: list[dict[str, object]] = []
    for config in sorted(by_config):
        config_rows = by_config[config]
        correct = sum(1 for row in config_rows if row["gold"] == row["predicted"])
        f1_values = [_f1_for_label(config_rows, label) for label in LABELS]
        coverage_values = [float(row.get("coverage") or 0) for row in config_rows]
        metrics.append(
            {
                "config": config,
                "n": len(config_rows),
                "accuracy": round(correct / len(config_rows), 12) if config_rows else 0.0,
                "macro_f1": round(sum(f1_values) / len(f1_values), 12),
                "coverage_mean": round(sum(coverage_values) / len(coverage_values), 12)
                if coverage_values
                else 0.0,
            }
        )
    return metrics


def _f1_for_label(rows: list[dict[str, str]], label: str) -> float:
    tp = sum(1 for row in rows if row["gold"] == label and row["predicted"] == label)
    fp = sum(1 for row in rows if row["gold"] != label and row["predicted"] == label)
    fn = sum(1 for row in rows if row["gold"] == label and row["predicted"] != label)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    return 2 * precision * recall / (precision + recall) if precision + recall else 0.0


def _confusion_by_config(rows: list[dict[str, str]]) -> list[dict[str, object]]:
    counts: Counter[tuple[str, str, str]] = Counter()
    for row in rows:
        counts[(row.get("config", ""), row["gold"], row["predicted"])] += 1
    return [
        {"config": config, "gold": gold, "predicted": predicted, "count": count}
        for (config, gold, predicted), count in sorted(counts.items())
    ]


def _write_summary(
    path: Path,
    metrics: list[dict[str, object]],
    has_gold: bool,
    gold_path: str | Path | None,
) -> None:
    lines = ["# Evaluation Summary", ""]
    if not has_gold:
        lines.extend([
            "No gold sentiment labels were available, so metric rows were not generated.",
            "",
        ])
    else:
        lines.append(f"Gold labels: `{gold_path}`" if gold_path else "Gold labels: scores file")
        lines.append("")
        lines.append("| Config | N | Accuracy | Macro-F1 | Coverage |")
        lines.append("| --- | ---: | ---: | ---: | ---: |")
        for row in metrics:
            lines.append(
                f"| {row['config']} | {row['n']} | {row['accuracy']} | "
                f"{row['macro_f1']} | {row['coverage_mean']} |"
            )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
