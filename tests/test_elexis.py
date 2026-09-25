from __future__ import annotations

from openpyxl import Workbook

from serbian_sentiment_wsd_eval.elexis import ElexisSense, ElexisSenseIndex, load_elexis_senses


def write_repo(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SrWSD-v2"
    sheet.append(["pos", "upos", "lemma", "senseID", "definition", "literals"])
    sheet.append(["N", "NOUN", "program", "ENG30-00551215-n", "javna priredba", "program"])
    sheet.append(["N", "NOUN", "program", "ENG30-06676416-n", "televizijski sadrzaj", "program"])
    sheet.append(["N", "NOUN", "program", "ENG30-00551215-n", "duplikat", "program"])
    sheet.append(["N", "NOUN", "emisija", "ENG30-06676416-n", "televizijski sadrzaj", "emisija"])
    sheet.append(["V", "VERB", "raditi", "ENG30-02413480-v", "", "raditi"])

    extra = workbook.create_sheet("SrWSD-eliminisanPOS")
    extra.append(["pos", "upos", "lemma", "senseID", "definition", "literals"])
    extra.append(["ADV", "ADV", "pre svega", "CVMWE-0332", "uglavnom", "pre svega"])

    unsupported = workbook.create_sheet("privremeno")
    unsupported.append(["id", "Def baza"])
    unsupported.append(["X", "ignored"])

    uppercase = workbook.create_sheet("priprema")
    uppercase.append(["UPOS", "lemma", "senseID", "", "Akcija", "literals", "definition"])
    uppercase.append(["ADJ", "anoniman", "ENG30-00120574-a", "", "dodati", "anonimni", "nepoznat identitet"])
    workbook.save(path)


def test_load_elexis_senses_defaults_to_srwsd_v2_and_deduplicates_ids(tmp_path):
    repo = tmp_path / "repo.xlsx"
    write_repo(repo)

    senses = load_elexis_senses(repo)

    assert [sense.sense_id for sense in senses] == [
        "ENG30-00551215-n",
        "ENG30-06676416-n",
        "ENG30-06676416-n",
    ]
    assert [sense.lemma for sense in senses] == ["program", "program", "emisija"]
    assert senses[0].definition == "javna priredba"


def test_load_elexis_senses_can_include_extra_supported_sheets(tmp_path):
    repo = tmp_path / "repo.xlsx"
    write_repo(repo)

    senses = load_elexis_senses(repo, sheets=["SrWSD-v2", "SrWSD-eliminisanPOS"])

    assert [sense.sense_id for sense in senses] == [
        "ENG30-00551215-n",
        "ENG30-06676416-n",
        "ENG30-06676416-n",
        "CVMWE-0332",
    ]


def test_load_elexis_senses_matches_headers_case_insensitively(tmp_path):
    repo = tmp_path / "repo.xlsx"
    write_repo(repo)

    senses = load_elexis_senses(repo, sheets=["priprema"])

    assert senses == [
        ElexisSense("ENG30-00120574-a", "anoniman", "ADJ", "nepoznat identitet")
    ]


def test_elexis_index_preserves_candidate_order_by_lemma_and_upos(tmp_path):
    repo = tmp_path / "repo.xlsx"
    write_repo(repo)

    index = ElexisSenseIndex(load_elexis_senses(repo))

    assert [sense.sense_id for sense in index.candidates_for_token("program", "NOUN")] == [
        "ENG30-00551215-n",
        "ENG30-06676416-n",
    ]
    assert [sense.sense_id for sense in index.candidates_for_token("emisija", "NOUN")] == [
        "ENG30-06676416-n",
    ]
    assert index.candidates_for_token("program", "VERB") == []
